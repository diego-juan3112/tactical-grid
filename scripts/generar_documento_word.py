"""Genera la version Word (.docx) del documento tecnico con el diseno editorial de DESIGN.md (E11.2).

La fuente es docs/documento-tecnico.md; el .docx se regenera cada vez que esta cambia. Requiere las dependencias
opcionales del grupo ``documentos`` (incluyen el ejecutable de pandoc), que no hacen falta para ejecutar el juego:
    python -m pip install -e ".[documentos]"
    python scripts/generar_documento_word.py            # solo el .docx
    python scripts/generar_documento_word.py --word     # ademas actualiza el indice con Word y exporta el PDF
"""

import argparse
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ_REPOSITORIO = Path(__file__).resolve().parents[1]
FUENTE = RAIZ_REPOSITORIO / "docs" / "documento-tecnico.md"
DESTINO = RAIZ_REPOSITORIO / "docs" / "documento-tecnico.docx"
URL_REPOSITORIO = "https://github.com/diego-juan3112/tactical-grid/blob/main/"
ENLACE_RELATIVO = re.compile(r"\]\((?!https?://|#)([^)\s]+)\)")
PIE_DE_PAGINA = "TacticalGrid · Documento técnico"

# Tokens de DESIGN.md, seccion 2. Ningun otro color puede aparecer en el documento.
PALETA = {
    "titulo_portada": "27272A",
    "titulo": "1E293B",
    "texto": "334155",
    "secundario": "64748B",
    "borde_fuerte": "CBD5E1",
    "borde": "E2E8F0",
    "nota_borde": "94A3B8",
    "superficie": "F8FAFC",
    "fondo": "FFFFFF",
}
TEXTO, TITULOS, CODIGO = "Segoe UI", "Segoe UI Semibold", "Consolas"

# Estilos de pandoc, DESIGN.md seccion 3: fuente, tamano, color, negrita, antes, despues e interlineado.
ESTILOS_DE_PARRAFO = {
    "Normal": (TEXTO, 10, "texto", False, 0, 8, 1.4),
    "Body Text": (TEXTO, 10, "texto", False, 0, 8, 1.4),
    "First Paragraph": (TEXTO, 10, "texto", False, 0, 8, 1.4),
    "Compact": (TEXTO, 10, "texto", False, 0, 4, 1.35),
    "Title": (TITULOS, 26, "titulo_portada", False, 96, 4, 1.1),
    "Subtitle": (TEXTO, 12, "secundario", False, 0, 24, 1.2),
    "Author": (TEXTO, 10.5, "texto", False, 0, 2, 1.2),
    "Date": (TEXTO, 9, "secundario", False, 18, 0, 1.2),
    "Heading 1": (TITULOS, 18, "titulo", False, 24, 8, 1.15),
    "Heading 2": (TITULOS, 13, "titulo", False, 18, 6, 1.2),
    "Heading 3": (TITULOS, 11, "titulo", False, 12, 4, 1.2),
    "Heading 4": (TITULOS, 10, "titulo", False, 12, 4, 1.2),
    "Block Text": (TEXTO, 10, "texto", False, 6, 10, 1.4),
    "Source Code": (CODIGO, 8.5, "texto", False, 6, 10, 1.25),
    "TOC Heading": (TITULOS, 13, "secundario", False, 0, 12, 1.2),
    "toc 1": (TITULOS, 10, "titulo", False, 5, 1, 1.15),
    "toc 2": (TEXTO, 9.5, "texto", False, 0, 1, 1.15),
    "Caption": (TEXTO, 8.5, "secundario", False, 4, 8, 1.2),
    "Table Caption": (TEXTO, 8.5, "secundario", False, 4, 8, 1.2),
}
ESTILOS_DE_CARACTER = {
    "Verbatim Char": (CODIGO, 8.5, "texto"),
    "Hyperlink": (TEXTO, None, "texto"),
}
COLORES_DEL_TEMA = {
    "dk2": "titulo", "lt2": "superficie", "accent1": "titulo", "accent2": "texto", "accent3": "secundario",
    "accent4": "nota_borde", "accent5": "borde_fuerte", "accent6": "borde", "hlink": "texto", "folHlink": "texto",
}


def preparar_markdown(texto: str) -> str:
    """Quita el titulo de nivel 1 y convierte los enlaces relativos en enlaces al repositorio.

    Proposito: adaptar la fuente Markdown al Word: el titulo sale de los metadatos YAML para la portada, y un
    enlace relativo a otro archivo no funciona fuera del repositorio.
    Precondiciones: ``texto`` puede empezar con un bloque YAML; el titulo es la primera linea que empieza con
    ``# `` fuera de ese bloque.
    Postcondiciones: devuelve el texto sin esa linea y con cada enlace relativo apuntando al archivo en la rama
    main de GitHub; el bloque YAML y el resto del contenido no cambian. No modifica el archivo fuente.
    Complejidad: O(n) temporal y espacial, donde n es la longitud del texto.
    Uso de IA: Si.
    Intervencion de IA: Claude (Opus 5.5) propuso esta funcion y su documentacion en E11.2.
    Validacion del estudiante: tests/test_documento_word.py comprueba que el .docx no repite el titulo y que sus
    enlaces son absolutos.
    """
    texto = re.sub(r"^# .*\n", "", texto, count=1, flags=re.MULTILINE)

    def absolutizar(coincidencia: re.Match[str]) -> str:
        ruta = (FUENTE.parent / coincidencia.group(1)).resolve().relative_to(RAIZ_REPOSITORIO)
        return f"]({URL_REPOSITORIO}{ruta.as_posix()})"

    return ENLACE_RELATIVO.sub(absolutizar, texto)


def _color(nombre: str) -> RGBColor:
    return RGBColor.from_string(PALETA[nombre])


def _quitar_atributos_de_tema(elemento) -> None:
    for atributo in list(elemento.attrib):
        if "Theme" in atributo or "theme" in atributo:
            del elemento.attrib[atributo]


def _fuente_y_color(estilo, fuente: str, tamano: float | None, color: str) -> None:
    propiedades = estilo.element.get_or_add_rPr()
    fuentes = propiedades.find(qn("w:rFonts"))
    if fuentes is None:
        fuentes = OxmlElement("w:rFonts")
        propiedades.insert(0, fuentes)
    _quitar_atributos_de_tema(fuentes)
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fuentes.set(qn(atributo), fuente)
    if tamano is not None:
        estilo.font.size = Pt(tamano)
    estilo.font.color.rgb = _color(color)
    _quitar_atributos_de_tema(propiedades.find(qn("w:color")))
    estilo.font.italic = False


def _borde(nombre: str, grosor_octavos: int, color: str, espacio: int = 0):
    borde = OxmlElement(f"w:{nombre}")
    borde.set(qn("w:val"), "single")
    borde.set(qn("w:sz"), str(grosor_octavos))
    borde.set(qn("w:space"), str(espacio))
    borde.set(qn("w:color"), PALETA[color])
    return borde


def _sombra(color: str):
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), PALETA[color])
    return sombra


def _reemplazar(padre, hijo) -> None:
    anterior = padre.find(hijo.tag)
    if anterior is not None:
        padre.remove(anterior)
    padre.append(hijo)


def _obtener_estilo(documento, nombre: str, tipo):
    for estilo in documento.styles:
        if estilo.name.lower() == nombre.lower():
            return estilo
    estilo = documento.styles.add_style(nombre, tipo)
    if nombre.startswith("toc "):
        # Word solo usa en el indice sus estilos integrados "toc N" (id TOCN), no uno personalizado homonimo.
        estilo.element.set(qn("w:styleId"), "TOC" + nombre.removeprefix("toc "))
        estilo.element.attrib.pop(qn("w:customStyle"), None)
    return estilo


def definir_estilos(documento) -> None:
    """Configura los estilos de pandoc con la paleta, la tipografia y los componentes de DESIGN.md.

    Proposito: concentrar el diseno en los estilos del documento de referencia, para que pandoc lo aplique a todo
    el contenido sin formato directo.
    Precondiciones: ``documento`` es el reference.docx predeterminado de pandoc abierto con python-docx.
    Postcondiciones: titulos, cuerpo, listas, notas, codigo, enlaces e indice usan solo fuentes y colores de
    DESIGN.md, sin colores de tema ni cursivas; las notas tienen barra izquierda y fondo, el codigo un marco fino,
    los encabezados quedan unidos al parrafo siguiente y los enlaces van subrayados sin azul.
    Complejidad: O(s) temporal, donde s es el numero de estilos configurados, y O(1) espacial adicional.
    Uso de IA: Si.
    Intervencion de IA: Claude (Opus 5.5) propuso esta funcion y su documentacion en E11.2.
    Validacion del estudiante: tests/test_documento_word.py comprueba que ningun color fuera de la paleta aparece
    en estilos ni contenido, y la revision visual se hizo sobre el PDF exportado por Word.
    """
    for nombre, (fuente, tamano, color, negrita, antes, despues, interlineado) in ESTILOS_DE_PARRAFO.items():
        estilo = _obtener_estilo(documento, nombre, WD_STYLE_TYPE.PARAGRAPH)
        _fuente_y_color(estilo, fuente, tamano, color)
        estilo.font.bold = negrita
        formato = estilo.paragraph_format
        formato.space_before, formato.space_after = Pt(antes), Pt(despues)
        formato.line_spacing = interlineado
        formato.widow_control = True
        if nombre.startswith("Heading") or nombre == "TOC Heading":
            formato.keep_with_next = True
    for nombre, (fuente, tamano, color) in ESTILOS_DE_CARACTER.items():
        _fuente_y_color(_obtener_estilo(documento, nombre, WD_STYLE_TYPE.CHARACTER), fuente, tamano, color)
    _obtener_estilo(documento, "Hyperlink", WD_STYLE_TYPE.CHARACTER).font.underline = True
    _obtener_estilo(documento, "toc 2", WD_STYLE_TYPE.PARAGRAPH).paragraph_format.left_indent = Cm(0.4)

    nota = _obtener_estilo(documento, "Block Text", WD_STYLE_TYPE.PARAGRAPH)
    nota.paragraph_format.left_indent = Cm(0.4)
    nota.paragraph_format.right_indent = Cm(0)
    propiedades = nota.element.get_or_add_pPr()
    bordes = OxmlElement("w:pBdr")
    bordes.append(_borde("left", 24, "nota_borde", espacio=8))
    _reemplazar(propiedades, bordes)
    _reemplazar(propiedades, _sombra("superficie"))

    codigo = _obtener_estilo(documento, "Source Code", WD_STYLE_TYPE.PARAGRAPH)
    propiedades = codigo.element.get_or_add_pPr()
    bordes = OxmlElement("w:pBdr")
    for lado in ("top", "left", "bottom", "right"):
        bordes.append(_borde(lado, 4, "borde", espacio=4))
    _reemplazar(propiedades, bordes)
    _reemplazar(propiedades, _sombra("superficie"))
    codigo.paragraph_format.left_indent = Cm(0.15)
    codigo.paragraph_format.right_indent = Cm(0.15)
    _limpiar_restos_del_tema(documento.styles.element)


def _limpiar_restos_del_tema(estilos) -> None:
    # Barrido final: estilos de pandoc no configurados y los predeterminados aun pueden heredar fuentes o colores
    # del tema; se fijan a Segoe UI y al color de texto, sin tocar los colores que ya son de la paleta.
    for fuentes in estilos.iter(qn("w:rFonts")):
        if any("Theme" in atributo for atributo in fuentes.attrib):
            _quitar_atributos_de_tema(fuentes)
            for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                fuentes.set(qn(atributo), fuentes.get(qn(atributo)) or TEXTO)
    for color in estilos.iter(qn("w:color")):
        _quitar_atributos_de_tema(color)
        if color.get(qn("w:val")) not in set(PALETA.values()) | {"auto"}:
            color.set(qn("w:val"), PALETA["texto"])


def aplicar_tema(ruta: Path) -> None:
    """Reemplaza los colores y fuentes del tema de Word por los de DESIGN.md.

    Proposito: eliminar el azul de raiz: pandoc y Word tinen titulos y enlaces con ``accent1`` y ``hlink`` del tema
    aunque los estilos no lo pidan.
    Precondiciones: ``ruta`` es un .docx con ``word/theme/theme1.xml``.
    Postcondiciones: los colores ``dk2``, ``lt2``, ``accent1`` a ``accent6``, ``hlink`` y ``folHlink`` del tema son
    de la paleta, las fuentes mayor y menor son Segoe UI, y el resto del paquete queda identico.
    Complejidad: O(z) temporal y espacial, donde z es el tamano descomprimido del paquete.
    Uso de IA: Si.
    Intervencion de IA: Claude (Opus 5.5) propuso esta funcion y su documentacion en E11.2.
    Validacion del estudiante: tests/test_documento_word.py comprueba que el tema no conserva ningun azul de Word.
    """
    with zipfile.ZipFile(ruta) as paquete:
        contenido = {nombre: paquete.read(nombre) for nombre in paquete.namelist()}
    tema = contenido["word/theme/theme1.xml"].decode("utf-8")
    for clave, token in COLORES_DEL_TEMA.items():
        tema = re.sub(rf"(<a:{clave}>)(.*?)(</a:{clave}>)",
                      rf'\g<1><a:srgbClr val="{PALETA[token]}"/>\g<3>', tema, flags=re.DOTALL)
    tema = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*', rf"\g<1>{TEXTO}", tema)
    contenido["word/theme/theme1.xml"] = tema.encode("utf-8")
    with zipfile.ZipFile(ruta, "w", zipfile.ZIP_DEFLATED) as paquete:
        for nombre, datos in contenido.items():
            paquete.writestr(nombre, datos)


def crear_referencia(destino: Path) -> Path:
    """Crea el reference.docx que pandoc usa como plantilla de estilos.

    Proposito: partir del documento de referencia oficial de pandoc, para no omitir ningun estilo que pandoc use,
    y aplicarle el diseno.
    Precondiciones: pypandoc_binary esta instalado; ``destino`` es una ruta escribible.
    Postcondiciones: ``destino`` contiene el reference.docx con los estilos de ``definir_estilos`` y el tema de
    ``aplicar_tema``; devuelve esa ruta.
    Complejidad: O(s + z), por los estilos y el tamano del paquete.
    Uso de IA: Si.
    Intervencion de IA: Claude (Opus 5.5) propuso esta funcion y su documentacion en E11.2.
    Validacion del estudiante: tests/test_documento_word.py genera el documento completo con esta referencia.
    """
    subprocess.run([pypandoc.get_pandoc_path(), "-o", str(destino), "--print-default-data-file",
                    "reference.docx"], check=True)
    documento = Document(str(destino))
    definir_estilos(documento)
    documento.save(str(destino))
    aplicar_tema(destino)
    return destino


def _salto_de_pagina():
    parrafo = OxmlElement("w:p")
    fragmento = OxmlElement("w:r")
    salto = OxmlElement("w:br")
    salto.set(qn("w:type"), "page")
    fragmento.append(salto)
    parrafo.append(fragmento)
    return parrafo


def _formatear_tabla(tabla) -> None:
    propiedades = tabla._tbl.tblPr
    ancho = OxmlElement("w:tblW")
    ancho.set(qn("w:w"), "5000")
    ancho.set(qn("w:type"), "pct")
    _reemplazar(propiedades, ancho)
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "bottom", "insideH"):
        bordes.append(_borde(lado, 4, "borde"))
    for lado in ("left", "right", "insideV"):
        sin_borde = OxmlElement(f"w:{lado}")
        sin_borde.set(qn("w:val"), "nil")
        bordes.append(sin_borde)
    _reemplazar(propiedades, bordes)
    margenes = OxmlElement("w:tblCellMar")
    for lado, valor in (("top", 60), ("bottom", 60), ("left", 100), ("right", 100)):
        margen = OxmlElement(f"w:{lado}")
        margen.set(qn("w:w"), str(valor))
        margen.set(qn("w:type"), "dxa")
        margenes.append(margen)
    _reemplazar(propiedades, margenes)
    for indice, fila in enumerate(tabla.rows):
        propiedades_fila = fila._tr.get_or_add_trPr()
        propiedades_fila.append(OxmlElement("w:cantSplit"))
        if indice == 0:
            repetir = OxmlElement("w:tblHeader")
            repetir.set(qn("w:val"), "true")
            propiedades_fila.append(repetir)
        for celda in fila.cells:
            if indice == 0:
                propiedades_celda = celda._tc.get_or_add_tcPr()
                _reemplazar(propiedades_celda, _sombra("superficie"))
                bordes_celda = OxmlElement("w:tcBorders")
                bordes_celda.append(_borde("bottom", 8, "borde_fuerte"))
                _reemplazar(propiedades_celda, bordes_celda)
            for parrafo in celda.paragraphs:
                parrafo.paragraph_format.space_before = Pt(2)
                parrafo.paragraph_format.space_after = Pt(2)
                parrafo.paragraph_format.line_spacing = 1.2
                for fragmento in parrafo.runs:
                    fragmento.font.size = Pt(8.5)
                    if indice == 0:
                        fragmento.font.name = TITULOS
                        fragmento.font.color.rgb = _color("titulo")


def _formatear_divisor(parrafo) -> None:
    for hijo in list(parrafo._p):
        if hijo.tag != qn("w:pPr"):
            parrafo._p.remove(hijo)
    propiedades = parrafo._p.get_or_add_pPr()
    bordes = OxmlElement("w:pBdr")
    bordes.append(_borde("bottom", 4, "borde"))
    _reemplazar(propiedades, bordes)
    parrafo.paragraph_format.space_before = Pt(18)
    parrafo.paragraph_format.space_after = Pt(0)
    marca = OxmlElement("w:rPr")
    tamano = OxmlElement("w:sz")
    tamano.set(qn("w:val"), "4")
    marca.append(tamano)
    _reemplazar(propiedades, marca)


def _agregar_pie(seccion) -> None:
    pie = seccion.footer
    parrafo = pie.paragraphs[0] if pie.paragraphs else pie.add_paragraph()
    parrafo.text = ""
    ancho = seccion.page_width - seccion.left_margin - seccion.right_margin
    parrafo.paragraph_format.tab_stops.add_tab_stop(ancho, WD_TAB_ALIGNMENT.RIGHT)
    texto = parrafo.add_run(f"{PIE_DE_PAGINA}\t")
    numero = parrafo.add_run()
    for tipo, instruccion in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            campo = OxmlElement("w:fldChar")
            campo.set(qn("w:fldCharType"), tipo)
        else:
            campo = OxmlElement("w:instrText")
            campo.set(qn("xml:space"), "preserve")
            campo.text = instruccion
        numero._r.append(campo)
    for fragmento in (texto, numero):
        fragmento.font.name = TEXTO
        fragmento.font.size = Pt(8)
        fragmento.font.color.rgb = _color("secundario")


def dar_formato(ruta: Path) -> None:
    """Aplica al .docx generado lo que los estilos de pandoc no pueden expresar.

    Proposito: completar DESIGN.md: pagina Carta con margenes de 2,5 cm, pie con numero de pagina, portada e indice
    en paginas propias, tablas sin bordes verticales, divisores finos y etiquetas de nota en color de titulo.
    Precondiciones: ``ruta`` es un .docx generado por pandoc con la referencia de ``crear_referencia``.
    Postcondiciones: el contenido no cambia; cada tabla ocupa el ancho del texto, con lineas horizontales de la
    paleta, encabezado sombreado que se repite y filas que no se parten; cada regla horizontal es una linea de
    0,5 pt; la portada y el indice terminan con un salto de pagina.
    Complejidad: O(p + c) temporal, con p parrafos y c celdas, y O(1) espacial adicional.
    Uso de IA: Si.
    Intervencion de IA: Claude (Opus 5.5) propuso esta funcion y su documentacion en E11.2.
    Validacion del estudiante: tests/test_documento_word.py comprueba bordes, encabezados repetidos, divisores,
    portada, pie de pagina y paleta; la revision visual se hizo sobre el PDF exportado por Word.
    """
    documento = Document(str(ruta))
    for seccion in documento.sections:
        seccion.start_type = WD_SECTION.NEW_PAGE
        seccion.page_width, seccion.page_height = Cm(21.59), Cm(27.94)
        for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(seccion, lado, Cm(2.5))
        _agregar_pie(seccion)
    for tabla in documento.tables:
        _formatear_tabla(tabla)
    for parrafo in documento.paragraphs:
        if parrafo._p.xpath(".//*[local-name()='rect' and @*[local-name()='hr']='t']"):
            _formatear_divisor(parrafo)
        elif parrafo.style.name == "Block Text" and parrafo.runs and parrafo.runs[0].bold:
            parrafo.runs[0].font.color.rgb = _color("titulo")
    indice = documento.element.body.find(qn("w:sdt"))
    if indice is not None:
        indice.addprevious(_salto_de_pagina())
        indice.addnext(_salto_de_pagina())
    documento.save(str(ruta))


def actualizar_con_word(ruta: Path) -> Path:
    """Rellena el indice con Microsoft Word y exporta el PDF de entrega.

    Proposito: pandoc deja el indice como un campo vacio que solo Word sabe calcular; la entrega final es un PDF
    exportado por Word, que incrusta las fuentes (DESIGN.md, seccion 8).
    Precondiciones: Windows con Microsoft Word instalado; ``ruta`` es el .docx ya formateado.
    Postcondiciones: el .docx queda guardado con el indice y los numeros de pagina calculados, y junto a el se
    escribe el .pdf con el mismo nombre; devuelve la ruta del PDF.
    Complejidad: O(p) respecto de las paginas, dominada por la paginacion de Word.
    Uso de IA: Si.
    Intervencion de IA: Claude (Opus 5.5) propuso esta funcion y su documentacion en E11.2.
    Validacion del estudiante: se ejecuto en E11.2 y se revisaron todas las paginas del PDF resultante.
    """
    pdf = ruta.with_suffix(".pdf")
    guion = (
        "$ErrorActionPreference = 'Stop'; $w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"try {{ $d = $w.Documents.Open('{ruta}'); $d.TablesOfContents | ForEach-Object {{ $_.Update() }}; "
        f"$d.Fields.Update() | Out-Null; $d.Save(); $d.ExportAsFixedFormat('{pdf}', 17); "
        "try { $d.Close(0) } catch { } } finally { try { $w.Quit() } catch { } }"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", guion], check=True)
    return pdf


def generar(fuente: Path = FUENTE, destino: Path = DESTINO) -> Path:
    """Convierte el documento tecnico de Markdown a Word con el diseno de DESIGN.md.

    Proposito: producir el entregable en .docx sin mantener dos fuentes: el Markdown sigue siendo la referencia.
    Precondiciones: ``fuente`` existe y estan instaladas las dependencias del grupo ``documentos``.
    Postcondiciones: ``destino`` contiene portada (titulo, subtitulo, autores y fecha del bloque YAML), indice de
    contenido hasta el nivel 2 y todas las secciones, con los capitulos como encabezados de nivel 1; devuelve la
    ruta generada.
    Complejidad: O(n) respecto del tamano del documento, dominada por la conversion de pandoc.
    Uso de IA: Si.
    Intervencion de IA: Claude (Opus 5.5) propuso esta funcion y su documentacion en E11.2.
    Validacion del estudiante: tests/test_documento_word.py genera el documento y comprueba su estructura y su
    diseno; la revision visual se hizo sobre el PDF exportado por Word.
    """
    with tempfile.TemporaryDirectory() as carpeta:
        referencia = crear_referencia(Path(carpeta) / "referencia.docx")
        pypandoc.convert_text(
            preparar_markdown(fuente.read_text(encoding="utf-8")),
            "docx",
            format="markdown",
            outputfile=str(destino),
            extra_args=[f"--reference-doc={referencia}", "--toc", "--toc-depth=2", "--shift-heading-level-by=-1",
                        "--syntax-highlighting=none", "-M", "toc-title=Contenido", "-M", "lang=es-CO"],
        )
    dar_formato(destino)
    return destino


def main() -> None:
    analizador = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    analizador.add_argument("--salida", type=Path, default=DESTINO, help="ruta del .docx a generar")
    analizador.add_argument("--word", action="store_true",
                            help="actualiza el indice con Microsoft Word y exporta el PDF (solo Windows)")
    argumentos = analizador.parse_args()
    ruta = generar(destino=argumentos.salida.resolve())
    print(f"Documento generado: {ruta}")
    if argumentos.word:
        print(f"PDF exportado: {actualizar_con_word(ruta)}")


if __name__ == "__main__":
    main()
